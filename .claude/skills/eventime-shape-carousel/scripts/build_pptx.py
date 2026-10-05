#!/usr/bin/env python3
"""Bangun carousel PPTX yang bisa diedit dari deck.json (deliverable utama skill ini).

Pakai:  python3 build_pptx.py deck.json out/C05.pptx [--no-anim]

Isi file:
- Semua teks, kotak, kartu, dan ikon adalah objek native PowerPoint (bisa digeser, diganti, diwarnai).
- Latar grid miring + logo Shape + 3 titik ada di Slide Master/Layout "Shape Carousel", jadi
  ganti sekali berlaku di semua slide.
- Highlight kata ({{...}} biru, [[...]] merah) memakai highlight teks PowerPoint (Microsoft 365 / 2019+).
- Motion = animasi masuk bawaan PowerPoint (Float In, Zoom, Wipe, Fade), jalan otomatis saat slide tampil.
  Ekspor ke video: File > Export > Create a Video.
- Catatan visual/elemen/motion/aset per slide ada di Speaker Notes.
- Font: Poppins (instal dari assets/fonts-ttf/). Tanpa Poppins PowerPoint memakai font pengganti.
"""
import argparse, copy, json, math, re, sys
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from PIL import ImageFont
from svgpathtools import parse_path

HERE = Path(__file__).resolve().parent
ASSETS = HERE.parent / "assets"
PX = 9525
W, H = 1080, 1350
E = lambda v: int(round(v * PX))

NAVY, HLBLUE, TBOX, RED, LIME = "0B233D", "4B7B98", "3D7396", "AB2224", "C8F03C"
PILLAR = {"mind": "8B6CF0", "fitness": "C8F03C", "nutrition": "FFD23F",
          "health": "37D3C2", "longevity": "7CC7FF", "aesthetics": "FF9EC4"}
HL_IN_BOX = "214765"      # highlight di dalam kotak biru: navy .55 di atas #3D7396

# role: (family, bold, italic, ttf)
FONTS = {
    "xb": ("Poppins ExtraBold", False, False, "Poppins-ExtraBold.ttf"),
    "bd": ("Poppins", True, False, "Poppins-Bold.ttf"),
    "sb": ("Poppins SemiBold", False, False, "Poppins-SemiBold.ttf"),
    "md": ("Poppins Medium", False, False, "Poppins-Medium.ttf"),
    "rg": ("Poppins", False, False, "Poppins-Regular.ttf"),
    "mi": ("Poppins Medium", False, True, "Poppins-MediumItalic.ttf"),
}
_fc = {}


def pil_font(role, size):
    k = (role, round(size))
    if k not in _fc:
        _fc[k] = ImageFont.truetype(str(ASSETS / "fonts-ttf" / FONTS[role][3]), round(size))
    return _fc[k]


# ------------------------------------------------------------------ markup & measuring
def segments(text):
    """'a {{b}} c [[d]]' -> [('a ',None),('b','hl'),(' c ',None),('d','red')]"""
    out, pos = [], 0
    for m in re.finditer(r"\{\{(.+?)\}\}|\[\[(.+?)\]\]", text, re.S):
        if m.start() > pos:
            out.append((text[pos:m.start()], None))
        out.append((m.group(1) or m.group(2), "hl" if m.group(1) else "red"))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], None))
    return out


NB = " "


def wrap_lines(text, role, size, maxw):
    """Perkiraan jumlah baris hasil wrap (greedy) termasuk pad NBSP highlight."""
    f = pil_font(role, size)
    total = 0
    widest = 0
    for para in text.split("\n"):
        toks = []
        for seg, kind in segments(para):
            words = seg.split(" ")
            for i, w in enumerate(words):
                if w == "" and len(words) > 1:
                    continue
                t = w
                if kind and i == 0:
                    t = NB + t
                if kind and i == len(words) - 1:
                    t = t + NB
                toks.append(t)
        line, n = "", 1
        for t in toks:
            cand = (line + " " + t) if line else t
            if f.getlength(cand) <= maxw or not line:
                line = cand
            else:
                widest = max(widest, f.getlength(line))
                n += 1
                line = t
        widest = max(widest, f.getlength(line))
        total += n
    return total, widest


def text_width(text, role, size):
    return wrap_lines(text, role, size, 10 ** 6)[1]


# ------------------------------------------------------------------ low-level shape helpers
def rgb(h):
    return RGBColor.from_string(h)


def set_alpha(parent, val_pct):
    """tambah <a:alpha> ke srgbClr pertama di dalam parent (solidFill)"""
    clr = parent.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        for old in clr.findall(qn("a:alpha")):
            clr.remove(old)
        al = etree.SubElement(clr, qn("a:alpha"))
        al.set("val", str(int(val_pct * 1000)))


def no_effects(shape):
    """Buang referensi style tema (sumber bayangan/garis bawaan) dan matikan efek."""
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)
    spPr = el.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def style_line(shape, color, width_px, alpha=None, dash=None, cap_round=False):
    shape.line.color.rgb = rgb(color)
    shape.line.width = E(width_px)
    ln = shape._element.spPr.find(qn("a:ln"))
    if alpha is not None:
        set_alpha(ln, alpha)
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", dash)
    if cap_round:
        ln.set("cap", "rnd")
        for tag in ("a:round",):
            ln.append(etree.Element(qn(tag)))


def fix_ln_order(shape):
    """Susun ulang anak <a:ln> sesuai skema: fill, prstDash, join(round), headEnd, tailEnd."""
    ln = shape._element.spPr.find(qn("a:ln"))
    if ln is None:
        return
    order = ["noFill", "solidFill", "gradFill", "pattFill", "prstDash", "custDash", "round", "bevel", "miter", "headEnd", "tailEnd"]
    kids = list(ln)
    kids.sort(key=lambda e: order.index(etree.QName(e).localname) if etree.QName(e).localname in order else 99)
    for k in kids:
        ln.remove(k)
    for k in kids:
        ln.append(k)


class Ctx:
    def __init__(self, slide, anim_on):
        self.slide = slide
        self.anims = []
        self.anim_on = anim_on

    def anim(self, shape, kind, delay):
        if shape is not None and self.anim_on and kind:
            self.anims.append((shape, kind, delay))
        return shape


def add_box(ctx, x, y, w, h, fill=None, fill_alpha=None, line=None, line_w=0, line_alpha=None,
            radius=0, dash=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rot=0, name=None):
    shp = ctx.slide.shapes.add_shape(shape if radius else (MSO_SHAPE.RECTANGLE if shape == MSO_SHAPE.ROUNDED_RECTANGLE else shape),
                                     E(x), E(y), E(w), E(h))
    if radius and shp.adjustments:
        shp.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        if fill_alpha is not None:
            set_alpha(shp._element.spPr.find(qn("a:solidFill")), fill_alpha)
    else:
        shp.fill.background()
    if line:
        style_line(shp, line, line_w, line_alpha, dash)
        fix_ln_order(shp)
    else:
        shp.line.fill.background()
    no_effects(shp)
    if rot:
        shp.rotation = rot
    if name:
        shp.name = name
    return shp


def fill_text(tf, text, role, size, color, align, lh, hl, alpha=None, anchor="t", wrap=True):
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    p.line_spacing = Pt(size * lh * 0.75)
    fam, bold, italic, _ = FONTS[role]
    first = True
    for line in text.split("\n"):
        if not first:
            p.add_line_break()
        first = False
        for seg, kind in segments(line):
            t = seg
            if kind:
                t = NB + seg + NB
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = fam; f.bold = bold; f.italic = italic
            f.size = Pt(size * 0.75)
            f.color.rgb = rgb("FFFFFF" if kind == "red" else color)
            rPr = r._r.get_or_add_rPr()
            if alpha is not None:
                set_alpha(rPr.find(qn("a:solidFill")), alpha)
            if kind:
                hlt = etree.Element(qn("a:highlight"))
                c = etree.SubElement(hlt, qn("a:srgbClr")); c.set("val", RED if kind == "red" else hl)
                rPr.find(qn("a:solidFill")).addnext(hlt)


def add_text(ctx, x, y, w, h, text, role, size, color="FFFFFF", align="l", lh=1.2, hl=HLBLUE,
             alpha=None, anchor="t", rot=0, name=None):
    tb = ctx.slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    fill_text(tb.text_frame, text, role, size, color, align, lh, hl, alpha, anchor)
    if rot:
        tb.rotation = rot
    if name:
        tb.name = name
    return tb


def add_shape_text(shp, text, role, size, color, lh=1.1, hl=HLBLUE, align="c"):
    fill_text(shp.text_frame, text, role, size, color, align, lh, hl, anchor="m", wrap=False)


# ------------------------------------------------------------------ icons (vektor, bisa diedit-titik)
ICON_D = {
    "brain": ["M9 4a3 3 0 0 0-3 3 3 3 0 0 0-2 5 3 3 0 0 0 2 5 3 3 0 0 0 6 1V4a3 3 0 0 0-3 0zM15 4a3 3 0 0 1 3 3 3 3 0 0 1 2 5 3 3 0 0 1-2 5 3 3 0 0 1-6 1"],
    "dumbbell": ["M6 7v10M3 9v6M18 7v10M21 9v6M6 12h12"],
    "bowl": ["M3 12h18a9 9 0 0 1-18 0zM8 8c0-2 2-2 2-4M13 8c0-2 2-2 2-4"],
    "heart": ["M12 20s-8-5-8-11a4.5 4.5 0 0 1 8-2.5A4.5 4.5 0 0 1 20 9c0 6-8 11-8 11z"],
    "leaf": ["M5 19C5 9 11 4 20 4c0 9-5 15-15 15zM5 19l8-8"],
    "spark": ["M12 3l2.2 6.8L21 12l-6.8 2.2L12 21l-2.2-6.8L3 12l6.8-2.2z"],
    "clock": ["M3 12a9 9 0 1 0 18 0a9 9 0 1 0-18 0z", "M12 7v5l3 2"],
    "calendar": ["M6 5h12a3 3 0 0 1 3 3v10a3 3 0 0 1-3 3H6a3 3 0 0 1-3-3V8a3 3 0 0 1 3-3z", "M3 10h18M8 3v4M16 3v4"],
    "search": ["M4 10.5a6.5 6.5 0 1 0 13 0a6.5 6.5 0 1 0-13 0z", "M15.5 15.5L21 21"],
    "check": ["M4 12.5l5 5L20 6.5"],
    "pulse": ["M2 12h5l2-6 4 12 2-6h7"],
    "moon": ["M20 14.5A8.5 8.5 0 0 1 9.5 4 8.5 8.5 0 1 0 20 14.5z"],
    "drop": ["M12 3s6 6.5 6 11a6 6 0 0 1-12 0c0-4.5 6-11 6-11z"],
    "link": ["M10 14a4 4 0 0 0 5.7 0l3-3a4 4 0 0 0-5.7-5.7l-1 1M14 10a4 4 0 0 0-5.7 0l-3 3a4 4 0 0 0 5.7 5.7l1-1"],
    "arrow": ["M4 12h15M13 5l7 7-7 7"],
}


def draw_polylines(ctx, polys, color, width_px, name, cap_round=True):
    """polys: list of list of (x_px, y_px). Satu freeform dengan beberapa kontur."""
    first = polys[0][0]
    fb = ctx.slide.shapes.build_freeform(E(first[0]), E(first[1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in polys[0][1:]], close=False)
    for pl in polys[1:]:
        fb.move_to(E(pl[0][0]), E(pl[0][1]))
        fb.add_line_segments([(E(x), E(y)) for x, y in pl[1:]], close=False)
    shp = fb.convert_to_shape(0, 0)
    shp.fill.background()
    style_line(shp, color, width_px, cap_round=cap_round)
    fix_ln_order(shp)
    no_effects(shp)
    shp.name = name
    return shp


def add_icon(ctx, name, x, y, size, color=NAVY, stroke=2.0):
    s = size / 24.0
    polys = []
    for d in ICON_D[name]:
        path = parse_path(d)
        for sub in path.continuous_subpaths():
            pts = []
            for seg in sub:
                n = max(4, int(seg.length() / 0.3))
                for i in range(n + 1):
                    z = seg.point(i / n)
                    pts.append((x + z.real * s, y + z.imag * s))
            polys.append(pts)
    return draw_polylines(ctx, polys, color, stroke * s, f"Ikon {name}")


# ------------------------------------------------------------------ komponen & layout vertikal
class Comp:
    def __init__(self, h, fn):
        self.h, self.fn = h, fn

    def draw(self, y, x0, w, align):
        return self.fn(y, x0, w, align)


def place_x(x0, w, cw, align):
    return x0 if align == "l" else (x0 + (w - cw) / 2 if align == "c" else x0 + w - cw)


def T(ctx, text, role, size, lh=1.2, color="FFFFFF", hl=HLBLUE, anim=None, delay=0, alpha=None, width=None, maxw=920, align=None, rot=0):
    """Teks ber-wrap. Tinggi dihitung dari lebar (width bila tetap, selain itu maxw = lebar wadah)."""
    n, _ = wrap_lines(text, role, size, (width or maxw) * 0.96)
    h = n * size * lh

    def fn(y, x0, w, al):
        ww = width or w
        tb = add_text(ctx, place_x(x0, w, ww, al), y, ww, h, text, role, size, color, align or al, lh, hl, alpha, rot=rot)
        ctx.anim(tb, anim, delay)
        return tb
    return Comp(h, fn)


def Pill(ctx, text, role, size, fill=None, color=NAVY, padx=30, pady=8, line=None, line_w=0, anim=None, delay=0, lh=1.15, rot=0):
    tw = text_width(text, role, size)
    pw, ph = tw + 2 * padx, size * lh + 2 * pady

    def fn(y, x0, w, al):
        xx = place_x(x0, w, pw, al)
        shp = add_box(ctx, xx, y, pw, ph, fill=fill, line=line, line_w=line_w, radius=min(ph / 2, 24) if fill else ph / 2, rot=rot)
        if not fill:
            shp.adjustments[0] = 0.5
        add_shape_text(shp, text, role, size, color, lh)
        ctx.anim(shp, anim, delay)
        return shp
    return Comp(ph, fn)


def IconC(ctx, name, d=150, fill="FFFFFF", color=NAVY, anim="pop", delay=0.1):
    def fn(y, x0, w, al):
        xx = place_x(x0, w, d, al)
        c = add_box(ctx, xx, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL, radius=0, name="Bulatan ikon")
        ic = add_icon(ctx, name, xx + d * 0.28, y + d * 0.28, d * 0.44, color, 2.0)
        ctx.anim(c, anim, delay)
        ctx.anim(ic, anim, delay)
        return c
    return Comp(d, fn)


def Asset(ctx, text, h, anim="rise", delay=1.0, fill=(255, 0.05)):
    def fn(y, x0, w, al):
        b = add_box(ctx, x0, y, w, h, fill="FFFFFF", fill_alpha=5, line="FFFFFF", line_w=4, line_alpha=50, radius=40, dash="dash", name="Placeholder aset")
        tf = b.text_frame
        fill_text(tf, "ASET\n" + text, "md", 28, "FFFFFF", "c", 1.3, HLBLUE, alpha=90, anchor="m")
        ctx.anim(b, anim, delay)
        return b
    return Comp(h, fn)


def Spacer(h):
    return Comp(h, lambda y, x0, w, al: None)


def stack(comps, gap, top, bottom, x0=80, w=920, align="c", valign="c"):
    total = sum(c.h for c in comps) + gap * (len(comps) - 1)
    y = max(top, top + (bottom - top - total) / 2) if valign == "c" else top
    for c in comps:
        c.draw(y, x0, w, align)
        y += c.h + gap
    return total


def BoxC(ctx, children, pad_x=52, pad_y=48, gap=26, fill=TBOX, radius=40, align="l", anim="rise", delay=0.3,
         line=None, line_w=0, line_alpha=None, fill_alpha=None):
    ih = sum(c.h for c in children) + gap * (len(children) - 1)
    h = ih + 2 * pad_y

    def fn(y, x0, w, al):
        bg = add_box(ctx, x0, y, w, h, fill=fill, fill_alpha=fill_alpha, line=line, line_w=line_w, line_alpha=line_alpha, radius=radius, name="Kotak")
        ctx.anim(bg, anim, delay)
        yy = y + pad_y
        for c in children:
            c.draw(yy, x0 + pad_x, w - 2 * pad_x, align)
            yy += c.h + gap
        return bg
    return Comp(h, fn)


# ------------------------------------------------------------------ animasi native PowerPoint
PNS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _set_vis(spid, i):
    return (f'<p:set><p:cBhvr><p:cTn id="{i}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="visible"/></p:to></p:set>')


def _fade(spid, i, dur):
    return (f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{i}" dur="{dur}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')


def _val(v):
    return f'<p:val><p:fltVal val="{v}"/></p:val>' if isinstance(v, (int, float)) else f'<p:val><p:strVal val="{v}"/></p:val>'


def _attr(spid, i, attr, v0, v1, dur):
    return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{i}" dur="{dur}" fill="hold"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:tavLst><p:tav tm="0">{_val(v0)}</p:tav><p:tav tm="100000">{_val(v1)}</p:tav></p:tavLst></p:anim>')


def effect_children(kind, spid, nid, dur):
    """-> (presetID, subtype, children xml)"""
    vis = _set_vis(spid, nid())
    if kind == "fade":
        return 10, 0, vis + _fade(spid, nid(), dur)
    if kind in ("rise",):
        return 42, 0, (vis + _fade(spid, nid(), dur) +
                       _attr(spid, nid(), "ppt_y", "#ppt_y+.04", "#ppt_y", dur))
    if kind in ("pop", "stamp"):
        return 53, 16, (vis + _attr(spid, nid(), "ppt_w", 0, "#ppt_w", dur) +
                        _attr(spid, nid(), "ppt_h", 0, "#ppt_h", dur) + _fade(spid, nid(), dur))
    if kind == "wipe":
        return 22, 8, vis + (f'<p:animEffect transition="in" filter="wipe(left)"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}"/>'
                              f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
    if kind == "slide-l":
        return 2, 8, (vis + _attr(spid, nid(), "ppt_x", "#ppt_x-.1", "#ppt_x", dur) + _fade(spid, nid(), dur))
    if kind == "slide-r":
        return 2, 2, (vis + _attr(spid, nid(), "ppt_x", "#ppt_x+.1", "#ppt_x", dur) + _fade(spid, nid(), dur))
    return 10, 0, vis + _fade(spid, nid(), dur)


def attach_animations(slide, anims):
    if not anims:
        return
    counter = [4]

    def nid():
        counter[0] += 1
        return counter[0]
    anims = sorted(anims, key=lambda a: a[2])
    nodes, bld = [], []
    for idx, (shape, kind, delay) in enumerate(anims):
        spid = shape.shape_id
        is_sp = shape._element.tag == qn("p:sp")
        dur = 700 if kind in ("rise", "slide-l", "slide-r", "wipe") else 500
        cid = nid()
        pid, sub, ch = effect_children(kind, spid, nid, dur)
        grp = ' grpId="0"' if is_sp else ""
        node = "afterEffect" if idx == 0 else "withEffect"
        nodes.append(f'<p:par><p:cTn id="{cid}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold"{grp} nodeType="{node}">'
                     f'<p:stCondLst><p:cond delay="{int(delay * 1000)}"/></p:stCondLst><p:childTnLst>{ch}</p:childTnLst></p:cTn></p:par>')
        if is_sp:
            bld.append(f'<p:bldP spid="{spid}" grpId="0" animBg="1"/>')
    xml = (f'<p:timing xmlns:p="{PNS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
           f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(nodes)}</p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{"".join(bld)}</p:bldLst></p:timing>')
    el = etree.fromstring(xml)
    sld = slide._element
    for old in sld.findall(qn("p:timing")):
        sld.remove(old)
    ext = sld.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(el)
    else:
        sld.append(el)


# ------------------------------------------------------------------ master/layout: latar, logo, titik
def build_layout(prs):
    layout = prs.slide_layouts[6]
    layout.name = "Shape Carousel"
    tree = layout.shapes._spTree
    for sp in list(tree):
        if sp.tag in (qn("p:sp"), qn("p:pic")):
            tree.remove(sp)

    def pic(img, x, y, w, h, name, sid):
        _, rid = layout.part.get_or_add_image_part(str(img))
        xml = (f'<p:pic xmlns:p="{PNS}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
               f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
               f'<p:nvPicPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/></p:nvPicPr>'
               f'<p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
               f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(h)}"/></a:xfrm>'
               f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
        tree.append(etree.fromstring(xml))

    def dot(x, y, d, sid, name):
        xml = (f'<p:sp xmlns:p="{PNS}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
               f'<p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
               f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(d)}" cy="{E(d)}"/></a:xfrm>'
               f'<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="67A1C4"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
        tree.append(etree.fromstring(xml))

    pic(ASSETS / "bg-grid.jpg", 0, 0, W, H, "Latar grid miring", 101)
    pic(ASSETS / "logo-shape.png", 442, 89, 201, 70, "Logo Shape", 102)
    for i, cx in enumerate((61.5, 107.5, 154.5)):
        dot(cx - 15.5, 77, 31, 103 + i, f"Titik {i + 1}")


# ------------------------------------------------------------------ pola slide
TOP, BOT = 210, 1100


def next_arrow(ctx, delay=1.6):
    ring = add_box(ctx, 853, 1123, 118, 118, line="FFFFFF", line_w=5, shape=MSO_SHAPE.OVAL, name="Panah next (cincin)")
    ring.fill.background()
    ar = draw_polylines(ctx, [[(881, 1182), (943, 1182)], [(924, 1163), (943, 1182), (924, 1201)]], "FFFFFF", 5, "Panah next")
    ctx.anim(ring, "fade", delay)
    ctx.anim(ar, "fade", delay)


def pat_cover(ctx, s):
    avail = BOT - TOP
    for tsize in (104, 96, 88, 80, 72):
        comps = []
        if s.get("icon"):
            comps.append(IconC(ctx, s["icon"], delay=0.1))
        if s.get("kicker"):
            comps.append(Pill(ctx, s["kicker"], "md", 34, fill=None, color="FFFFFF", line="FFFFFF", line_w=3, anim="rise", delay=0.2))
        comps.append(T(ctx, s["title"], "xb", tsize, 1.18, anim="rise", delay=0.4))
        if s.get("sub"):
            comps.append(T(ctx, s["sub"], "md", 44, 1.4, width=860, anim="rise", delay=0.8))
        if s.get("asset") or s.get("img"):
            comps.append(Asset(ctx, s.get("asset", ""), s.get("assetH", 300), delay=1.0))
        if sum(c.h for c in comps) + 36 * (len(comps) - 1) <= avail:
            break
    if s.get("ghost"):
        g = add_text(ctx, 0, 440, W, 340, s["ghost"], "xb", 300, "FFFFFF", "c", 1.0, alpha=7, name="Kata ghost")
        ctx.anim(g, "fade", 0.1)
    stack(comps, 36, TOP, BOT, align="c")


def pat_claim(ctx, s):
    kids = []
    if s.get("label"):
        kids.append(Pill(ctx, s["label"], "bd", 46, fill="FFFFFF", color=NAVY, padx=30, pady=8, anim="pop", delay=0.45))
    kids.append(T(ctx, s["body"], "md", s.get("size", 52), 1.4, hl=HL_IN_BOX, anim="rise", delay=0.6, maxw=696))
    inner = [BoxC(ctx, kids, gap=26, fill=TBOX, anim="rise", delay=0.3)]
    if s.get("asset") or s.get("img"):
        inner.append(Asset(ctx, s.get("asset", ""), s.get("assetH", 360), delay=0.9))
    card = BoxC(ctx, inner, pad_x=60, pad_y=56, gap=36, fill="FFFFFF", fill_alpha=8, line="FFFFFF", line_w=4, line_alpha=55, radius=64, anim="fade", delay=0.1)
    stack([card], 0, 210, 1090, align="l")


def pat_numbered(ctx, s):
    comps = [T(ctx, str(s["num"]), "xb", 260, 0.8, color="FFFFFF", alpha=24, anim="fade", delay=0.1),
             Pill(ctx, s.get("label") or f"Tanda #{s['num']}", "bd", 40, fill=HLBLUE, color="FFFFFF", padx=26, pady=8, anim="pop", delay=0.3, rot=-2),
             T(ctx, s["title"], "xb", 88, 1.18, anim="rise", delay=0.5)]
    if s.get("body"):
        comps.append(T(ctx, s["body"], "md", 44, 1.4, anim="rise", delay=0.8))
    if s.get("asset") or s.get("img"):
        comps.append(Asset(ctx, s.get("asset", ""), s.get("assetH", 260), delay=1.0))
    stack(comps, 36, TOP, BOT, align="l")


def pat_list(ctx, s):
    items = s["items"]
    title = T(ctx, s["title"], "xb", 76, 1.18, anim="rise", delay=0.1)
    row_h = 120
    rows = []
    for i, t in enumerate(items):
        def mk(t=t, i=i):
            def fn(y, x0, w, al):
                card = add_box(ctx, x0, y, w, row_h, fill="FFFFFF", fill_alpha=8, line="FFFFFF", line_w=4, line_alpha=55, radius=44, name=f"Baris {i + 1}")
                d = 76
                c = add_box(ctx, x0 + 34, y + (row_h - d) / 2, d, d, fill=HLBLUE, shape=MSO_SHAPE.OVAL, name="Bulatan centang")
                ck = add_icon(ctx, "check", x0 + 34 + d * 0.25, y + (row_h - d) / 2 + d * 0.25, d * 0.5, "FFFFFF", 2.6)
                tx = add_text(ctx, x0 + 34 + d + 28, y, w - (34 + d + 28 + 34), row_h, t, "md", 40, "FFFFFF", "l", 1.2, anchor="m", name="Teks baris")
                for sh in (card, c, ck, tx):
                    ctx.anim(sh, "slide-l", 0.35 + i * 0.25)
            return Comp(row_h, fn)
        rows.append(mk())
    stack([title] + rows, 24, 230, BOT, align="l", valign="c")


def pat_compare(ctx, s):
    comps = []
    if s.get("title"):
        comps.append(T(ctx, s["title"], "bd", 64, 1.2, anim="rise", delay=0.1, align="c"))
    L, R = s["left"], s["right"]
    cw = 446

    def col(d, red, anim, delay, fill, fa):
        kids = [Pill(ctx, d["label"], "bd", 36, fill=RED if red else "FFFFFF", color="FFFFFF" if red else NAVY, padx=26, pady=8, anim=anim, delay=delay),
                T(ctx, d["title"], "bd", 64, 1.2, anim=anim, delay=delay, width=cw - 88),
                T(ctx, d["body"], "md", 34, 1.4, anim=anim, delay=delay, width=cw - 88)]
        return kids

    lk, rk = col(L, False, "slide-l", 0.3, None, 8), col(R, bool(R.get("red")), "slide-r", 0.5, HLBLUE, 45)
    gap = 26
    hh = max(sum(c.h for c in lk) + gap * 2, sum(c.h for c in rk) + gap * 2) + 88

    def row_fn(y, x0, w, al):
        for idx, (kids, fill, fa, anim, delay) in enumerate(((lk, "FFFFFF", 8, "slide-l", 0.3), (rk, HLBLUE, 45, "slide-r", 0.5))):
            x = x0 + idx * (cw + 28)
            bg = add_box(ctx, x, y, cw, hh, fill=fill, fill_alpha=fa, line="FFFFFF", line_w=4, line_alpha=55, radius=64, name="Kartu")
            ctx.anim(bg, anim, delay)
            yy = y + 44
            for c in kids:
                c.draw(yy, x + 44, cw - 88, "l")
                yy += c.h + gap
    comps.append(Comp(hh, row_fn))
    if s.get("foot"):
        comps.append(T(ctx, s["foot"], "md", 44, 1.4, anim="rise", delay=1.0, align="c"))
    stack(comps, 36, TOP, BOT, align="l")


def pat_bars(ctx, s):
    comps = [T(ctx, s["title"], "bd", 64, 1.2, anim="rise", delay=0.1)]
    for i, it in enumerate(s["items"]):
        def mk(it=it, i=i):
            def fn(y, x0, w, al):
                d = 0.2 + i * 0.2
                lab = add_text(ctx, x0, y, w * 0.7, 44, it["label"], "md", 34, "FFFFFF", "l", 1.3)
                val = add_text(ctx, x0 + w * 0.7, y, w * 0.3, 44, str(it.get("text", it["value"])), "bd", 34, "FFFFFF", "r", 1.3)
                bar = add_box(ctx, x0, y + 56, w * it["value"] / 100.0, 84, fill=RED if it.get("red") else "FFFFFF", radius=24, name="Batang")
                for sh in (lab, val):
                    ctx.anim(sh, "fade", d)
                ctx.anim(bar, "wipe", d + 0.2)
            return Comp(140, fn)
        comps.append(mk())
    if s.get("caption"):
        comps.append(T(ctx, s["caption"], "md", 44, 1.4, anim="rise", delay=1.0))
    stack(comps, 36, TOP, BOT, align="l")


def pat_stat(ctx, s):
    comps = []
    if s.get("kicker"):
        comps.append(Pill(ctx, s["kicker"], "md", 34, fill=None, color="FFFFFF", line="FFFFFF", line_w=3, anim="rise", delay=0.1))
    comps.append(T(ctx, f"{s['number']}{s.get('suffix', '')}", "xb", 260, 1.0, anim="pop", delay=0.3))
    comps.append(T(ctx, s["label"], "bd", 64, 1.2, anim="rise", delay=0.7))
    if s.get("body"):
        comps.append(T(ctx, s["body"], "md", 44, 1.4, width=860, anim="rise", delay=1.0))
    stack(comps, 36, TOP, BOT, align="c")


def pat_search(ctx, s):
    bar_h = 124

    def bar(y, x0, w, al):
        b = add_box(ctx, x0, y, w, bar_h, fill="FFFFFF", radius=62, name="Search bar")
        ic = add_icon(ctx, "search", x0 + 40, y + (bar_h - 64) / 2, 64, NAVY, 2.4)
        tx = add_text(ctx, x0 + 40 + 64 + 26, y, w - 170, bar_h, s["query"], "sb", 48, NAVY, "l", 1.2, anchor="m", name="Query")
        for sh in (b, ic, tx):
            ctx.anim(sh, "rise", 0.1)
    comps = [Comp(bar_h, bar),
             BoxC(ctx, [T(ctx, s["body"], "md", 52, 1.4, hl=HL_IN_BOX, anim="rise", delay=0.8, maxw=816)], gap=0, anim="rise", delay=0.7)]
    if s.get("asset") or s.get("img"):
        comps.append(Asset(ctx, s.get("asset", ""), s.get("assetH", 300), delay=1.2))
    stack(comps, 36, TOP, BOT, align="l")


def pat_session(ctx, s):
    comps = []
    if s.get("bigLogo"):
        def lg(y, x0, w, al):
            _, rid = None, None
            p = ctx.slide.shapes.add_picture(str(ASSETS / "logo-shape.png"), E(x0 + (w - 560) / 2), E(y), E(560), E(195))
            p.name = "Logo Shape besar"
            ctx.anim(p, "pop", 0.1)
        comps.append(Comp(195, lg))
    comps.append(Pill(ctx, s.get("kicker", "Shape Expert Network"), "md", 34, fill=None, color="FFFFFF", line="FFFFFF", line_w=3, anim="fade", delay=0.1))
    comps.append(T(ctx, s["title"], "xb", 96, 1.18, anim="rise", delay=0.3))
    if s.get("sub"):
        comps.append(BoxC(ctx, [T(ctx, s["sub"], "md", 44, 1.4, hl=HL_IN_BOX, align="c", anim="rise", delay=0.7, maxw=816)], pad_x=52, pad_y=36, gap=0, anim="rise", delay=0.7))
    if s.get("date"):
        def dt(y, x0, w, al):
            tw = text_width(s["date"], "bd", 64)
            tot = 104 + 22 + tw
            xx = x0 + (w - tot) / 2
            c = add_box(ctx, xx, y, 104, 104, fill="FFFFFF", shape=MSO_SHAPE.OVAL, name="Bulatan kalender")
            ic = add_icon(ctx, "calendar", xx + 26, y + 26, 52, NAVY, 2.0)
            t = add_text(ctx, xx + 126, y, tw + 20, 104, s["date"], "bd", 64, "FFFFFF", "l", 1.2, anchor="m")
            for sh in (c, ic, t):
                ctx.anim(sh, "rise", 1.0)
        comps.append(Comp(104, dt))
    stack(comps, 36, TOP, BOT, align="c")


def pat_quote(ctx, s):
    comps = [T(ctx, "“", "xb", 150, 0.7, alpha=35, anim="pop", delay=0.1),
             BoxC(ctx, [T(ctx, s["quote"], "mi", 50, 1.4, hl=HL_IN_BOX, anim="rise", delay=0.5, maxw=816)], gap=0, anim="rise", delay=0.4),
             Pill(ctx, s["who"], "bd", 38, fill="FFFFFF", color=NAVY, padx=26, pady=8, anim="rise", delay=0.9)]
    if s.get("role"):
        comps.append(T(ctx, s["role"], "md", 34, 1.4, anim="rise", delay=1.0))
    stack(comps, 30, TOP, BOT, align="l")


def pat_logos(ctx, s):
    comps = [T(ctx, s["title"], "xb", 68, 1.18, anim="rise", delay=0.1, align="c")]
    names = s["logos"]
    cw, ch, g = 290, 170, 22
    rows = [names[i:i + 3] for i in range(0, len(names), 3)]

    def grid(y, x0, w, al):
        yy = y
        k = 0
        for r in rows:
            tw = len(r) * cw + (len(r) - 1) * g
            xx = x0 + (w - tw) / 2
            for n in r:
                b = add_box(ctx, xx, yy, cw, ch, fill="FFFFFF", line="9BB5C4", line_w=3, dash="dash", radius=36, name=f"Logo {n}")
                fill_text(b.text_frame, "LOGO\n" + n, "bd", 26, NAVY, "c", 1.3, HLBLUE, anchor="m")
                ctx.anim(b, "pop", 0.4 + k * 0.15)
                xx += cw + g
                k += 1
            yy += ch + g
    comps.append(Comp(len(rows) * ch + (len(rows) - 1) * g, grid))
    if s.get("note"):
        comps.append(T(ctx, s["note"], "md", 44, 1.4, anim="rise", delay=1.4, align="c"))
    stack(comps, 30, 240, BOT, align="c")


def pat_timeline(ctx, s):
    steps = s["steps"]
    comps = [T(ctx, s["title"], "xb", 78, 1.18, anim="rise", delay=0.1)]
    ch = 124

    def tl(y, x0, w, al):
        n = len(steps)
        total = n * ch + (n - 1) * 26
        line = add_box(ctx, x0 + 22, y + 20, 6, total - 40, fill="FFFFFF", fill_alpha=35, radius=3, name="Garis timeline")
        ctx.anim(line, "wipe", 0.3)
        yy = y
        for i, st in enumerate(steps):
            card = add_box(ctx, x0 + 70, yy, w - 70, ch, fill="FFFFFF", fill_alpha=8, line="FFFFFF", line_w=3, line_alpha=55, radius=36, name=f"Langkah {i + 1}")
            dot = add_box(ctx, x0 + 10, yy + ch / 2 - 17, 34, 34, fill="FFFFFF", line=HLBLUE, line_w=8, shape=MSO_SHAPE.OVAL, name="Titik")
            w1 = add_text(ctx, x0 + 70 + 34, yy + 18, w - 70 - 68, 36, st["when"], "md", 30, "FFFFFF", "l", 1.2, alpha=75)
            w2 = add_text(ctx, x0 + 70 + 34, yy + 54, w - 70 - 68, 56, st["what"], "bd", 40, "FFFFFF", "l", 1.2)
            for sh in (card, dot, w1, w2):
                ctx.anim(sh, "rise", 0.5 + i * 0.2)
            yy += ch + 26
    comps.append(Comp(len(steps) * ch + (len(steps) - 1) * 26, tl))
    stack(comps, 36, 230, BOT, align="l")


def pat_speaker(ctx, s):
    pc = PILLAR.get(s.get("pillar"), PILLAR["fitness"])
    right = s.get("align", "right") != "left"
    side_x = (W - 64) if right else 64
    al = "r" if right else "l"
    # foto / placeholder di sisi sebaliknya
    ph_w, ph_h = 580, 800
    px = 0 if right else W - ph_w
    if s.get("photo"):
        pic = ctx.slide.shapes.add_picture(str(Path(s["photo"])), E(px), E(H - ph_h), height=E(ph_h))
        pic.name = "Foto pembicara"
        ctx.anim(pic, "rise", 0.2)
    else:
        b = add_box(ctx, px, H - ph_h, ph_w, ph_h, fill="FFFFFF", fill_alpha=12, line="FFFFFF", line_w=4, line_alpha=40, dash="dash",
                    shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, radius=1, name="Placeholder foto pembicara")
        b.adjustments[0] = 0.5
        fill_text(b.text_frame, "FOTO CUTOUT\n" + s["name"], "md", 34, "FFFFFF", "c", 1.3, HLBLUE, alpha=85, anchor="m")
        ctx.anim(b, "rise", 0.2)
    # ikon + tag kategori
    d = 150
    ix = side_x - d if right else side_x
    circ = add_box(ctx, ix, 200, d, d, fill=pc, shape=MSO_SHAPE.OVAL, name="Bulatan ikon")
    ic = add_icon(ctx, s.get("icon", "spark"), ix + d * 0.28, 200 + d * 0.28, d * 0.44, NAVY, 2.0)
    ctx.anim(circ, "pop", 0.2); ctx.anim(ic, "pop", 0.2)
    tw = text_width(s["cat"], "xb", 104) + 68
    tx = side_x - tw if right else side_x
    tag = add_box(ctx, tx, 370, tw, 128, fill=pc, radius=14, rot=-3, name="Tag kategori")
    add_shape_text(tag, s["cat"], "xb", 104, NAVY, 1.0)
    ctx.anim(tag, "pop", 0.4)
    # nama & jabatan
    colw = 430
    nx = side_x - colw if right else side_x
    n_lines, _ = wrap_lines(s["name"], "xb", 40, colw - 52)
    nh = n_lines * 40 * 1.2 + 26
    nm = add_box(ctx, nx, 560, colw, nh, fill="FFFFFF", radius=14, rot=-1.5, name="Tag nama")
    fill_text(nm.text_frame, s["name"], "xb", 40, NAVY, "r" if right else "l", 1.2, HLBLUE, anchor="m")
    nm.text_frame.margin_left = nm.text_frame.margin_right = E(26)
    r_lines, _ = wrap_lines(s["role"], "xb", 38, colw - 52)
    rh = r_lines * 38 * 1.2 + 22
    rl = add_box(ctx, nx, 560 + nh + 8, colw, rh, fill="CFE6FF", radius=14, rot=-1, name="Tag jabatan")
    fill_text(rl.text_frame, s["role"], "xb", 38, NAVY, "r" if right else "l", 1.2, HLBLUE, anchor="m")
    rl.text_frame.margin_left = rl.text_frame.margin_right = E(26)
    ctx.anim(nm, "slide-r" if right else "slide-l", 0.7); ctx.anim(rl, "slide-r" if right else "slide-l", 0.85)
    yy = 560 + nh + 8 + rh + 28
    if s.get("quote"):
        q = f"“{s['quote']}”"
        n, _ = wrap_lines(q, "mi", 40, colw)
        qh = n * 40 * 1.35
        qt = add_text(ctx, nx, yy, colw, qh, q, "mi", 40, "FFFFFF", al, 1.35)
        sc = add_box(ctx, nx + (colw * 0.22 if right else 0), yy + qh + 8, colw * 0.78, 14, fill=LIME, radius=7, name="Scribble")
        ctx.anim(qt, "rise", 1.1); ctx.anim(sc, "wipe", 1.5)
        yy += qh + 40
    if s.get("date"):
        dw = text_width(s["date"], "xb", 64) + 68
        dx = side_x - dw if right else side_x
        st = add_box(ctx, dx, min(max(yy + 10, 960), 1110), dw, 100, fill=RED, radius=28, rot=-6, name="Sticker tanggal")
        add_shape_text(st, s["date"], "xb", 64, "FFFFFF", 1.0)
        ctx.anim(st, "pop", 1.3)
    # CTA footer
    ct = "Daftar via link di bio →"
    cw = text_width(ct, "xb", 34) + 72
    cx = 64 if right else W - 64 - cw
    cta = add_box(ctx, cx, 1230, cw, 70, fill=LIME, radius=35, name="CTA pill")
    add_shape_text(cta, ct, "xb", 34, NAVY, 1.0)
    ctx.anim(cta, "fade", 1.6)


def pat_cta(ctx, s):
    comps = []
    if s.get("icon"):
        comps.append(IconC(ctx, s["icon"], delay=0.1))
    comps.append(BoxC(ctx, [T(ctx, s["title"], "xb", 84, 1.18, hl=HL_IN_BOX, align="c", anim="rise", delay=0.4, maxw=816)], pad_x=52, pad_y=48, gap=0, anim="rise", delay=0.3))
    if s.get("sub"):
        comps.append(T(ctx, s["sub"], "bd", 56, 1.25, anim="rise", delay=0.8, align="c"))
    if s.get("handle"):
        comps.append(Pill(ctx, s["handle"], "xb", 44, fill=LIME, color=NAVY, padx=34, pady=10, anim="pop", delay=1.1))
    if s.get("asset") or s.get("img"):
        comps.append(Asset(ctx, s.get("asset", ""), s.get("assetH", 260), delay=1.2))
    stack(comps, 36, TOP, BOT, align="c")


PATTERNS = {"cover": pat_cover, "claim": pat_claim, "numbered": pat_numbered, "list": pat_list, "compare": pat_compare,
            "bars": pat_bars, "stat": pat_stat, "search": pat_search, "session": pat_session, "quote": pat_quote,
            "logos": pat_logos, "timeline": pat_timeline, "speaker": pat_speaker, "cta": pat_cta}


def build(deck, out, anim_on=True):
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    build_layout(prs)
    layout = prs.slide_layouts[6]
    meta, slides = deck.get("meta", {}), deck["slides"]
    prs.core_properties.title = meta.get("title", "Carousel")
    prs.core_properties.author = "Eventime x Shape Indonesia"
    for i, s in enumerate(slides):
        slide = prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        ctx = Ctx(slide, anim_on)
        PATTERNS[s["pattern"]](ctx, s)
        if s["pattern"] != "speaker" and i != len(slides) - 1:
            next_arrow(ctx)
        n = s.get("notes") or {}
        txt = "\n".join(f"{k.upper()}: {v}" for k, v in (("Pola", s["pattern"]), ("Visual", n.get("visual")), ("Elemen", n.get("element")),
                                                         ("Motion", n.get("motion")), ("Aset", n.get("asset"))) if v)
        slide.notes_slide.notes_text_frame.text = txt
        attach_animations(slide, ctx.anims)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return len(slides)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck"); ap.add_argument("out"); ap.add_argument("--no-anim", action="store_true")
    a = ap.parse_args()
    deck = json.loads(Path(a.deck).read_text())
    sys.path.insert(0, str(HERE))
    from build_preview import check
    errs, warns = check(deck)
    for w in warns: print("WARN ", w)
    for e in errs: print("ERROR", e)
    n = build(deck, a.out, not a.no_anim)
    print(f"OK  {a.out} ({n} slide)")
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
